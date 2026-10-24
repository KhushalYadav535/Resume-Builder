import openpyxl

wb1_path = r'd:\resume\UpRole_21_Resume_Batch_Extraction_Report.xlsx'
wb2_path = r'd:\resume\UpRole_21_Resume_Batch_Extraction_Report_v2.xlsx'

wb1 = openpyxl.load_workbook(wb1_path, data_only=True)
wb2 = openpyxl.load_workbook(wb2_path, data_only=True)

print("=== SHEET NAMES ===")
print("WB1 Sheets:", wb1.sheetnames)
print("WB2 Sheets:", wb2.sheetnames)

# Compare Sheet row counts
print("\n=== ROW COUNTS COMPARISON ===")
for sname in wb1.sheetnames:
    ws1 = wb1[sname]
    ws2 = wb2[sname]
    print(f"{sname:<30} | v1 Data Rows: {ws1.max_row - 4:<5} | v2 Data Rows: {ws2.max_row - 4:<5} | Diff: {ws2.max_row - ws1.max_row:+d}")

# Compare Executive_Summary row by row
ws1_sum = wb1['Executive_Summary']
ws2_sum = wb2['Executive_Summary']

def clean_int(val):
    if val is None:
        return 0
    s = str(val).replace('%', '').strip()
    try:
        return int(float(s))
    except:
        return 0

data1 = []
for r in range(1, ws1_sum.max_row + 1):
    val = str(ws1_sum.cell(row=r, column=1).value or '')
    if val.startswith('RES_'):
        data1.append({
            'id': val,
            'file': ws1_sum.cell(row=r, column=2).value,
            'name': ws1_sum.cell(row=r, column=5).value or '',
            'headline': ws1_sum.cell(row=r, column=6).value or '',
            'roles': clean_int(ws1_sum.cell(row=r, column=8).value),
            'resps': clean_int(ws1_sum.cell(row=r, column=9).value),
            'outcomes': clean_int(ws1_sum.cell(row=r, column=10).value),
            'metrics': clean_int(ws1_sum.cell(row=r, column=11).value),
            'skills': clean_int(ws1_sum.cell(row=r, column=12).value),
            'edu': clean_int(ws1_sum.cell(row=r, column=13).value),
            'certs': clean_int(ws1_sum.cell(row=r, column=14).value),
            'facts': clean_int(ws1_sum.cell(row=r, column=15).value),
            'score': clean_int(ws1_sum.cell(row=r, column=16).value),
            'status': str(ws1_sum.cell(row=r, column=17).value or ''),
        })

data2 = []
for r in range(1, ws2_sum.max_row + 1):
    val = str(ws2_sum.cell(row=r, column=1).value or '')
    if val.startswith('RES_'):
        data2.append({
            'id': val,
            'file': ws2_sum.cell(row=r, column=2).value,
            'name': ws2_sum.cell(row=r, column=5).value or '',
            'headline': ws2_sum.cell(row=r, column=6).value or '',
            'roles': clean_int(ws2_sum.cell(row=r, column=8).value),
            'resps': clean_int(ws2_sum.cell(row=r, column=9).value),
            'outcomes': clean_int(ws2_sum.cell(row=r, column=10).value),
            'metrics': clean_int(ws2_sum.cell(row=r, column=11).value),
            'skills': clean_int(ws2_sum.cell(row=r, column=12).value),
            'edu': clean_int(ws2_sum.cell(row=r, column=13).value),
            'certs': clean_int(ws2_sum.cell(row=r, column=14).value),
            'facts': clean_int(ws2_sum.cell(row=r, column=15).value),
            'score': clean_int(ws2_sum.cell(row=r, column=16).value),
            'status': str(ws2_sum.cell(row=r, column=17).value or ''),
        })

print("\n=== DETAILED RESUME-BY-RESUME COMPARISON TABLE ===")
print(f"{'ID':<7} | {'Candidate (v1 -> v2)':<35} | {'Roles':<10} | {'Resps':<10} | {'Score':<14} | {'Status (v1 -> v2)'}")
print("-" * 115)

for d1, d2 in zip(data1, data2):
    name_str = f"{d1['name'][:15]} -> {d2['name'][:15]}" if d1['name'] != d2['name'] else d1['name'][:25]
    roles_str = f"{d1['roles']} -> {d2['roles']}" if d1['roles'] != d2['roles'] else str(d1['roles'])
    resps_str = f"{d1['resps']} -> {d2['resps']}" if d1['resps'] != d2['resps'] else str(d1['resps'])
    score_str = f"{d1['score']}% -> {d2['score']}%" if d1['score'] != d2['score'] else f"{d1['score']}%"
    status_str = f"{d1['status']} -> {d2['status']}" if d1['status'] != d2['status'] else d1['status']
    print(f"{d1['id']:<7} | {name_str:<35} | {roles_str:<10} | {resps_str:<10} | {score_str:<14} | {status_str}")

print("\n=== AGGREGATE TOTALS (v1 vs v2) ===")
avg_score1 = sum(d['score'] for d in data1) / len(data1)
avg_score2 = sum(d['score'] for d in data2) / len(data2)
print(f"Average Accuracy Score:  {avg_score1:.1f}% -> {avg_score2:.1f}% (+{avg_score2 - avg_score1:.1f}% improvement)")

roles1 = sum(d['roles'] for d in data1)
roles2 = sum(d['roles'] for d in data2)
print(f"Total Roles Extracted:   {roles1} -> {roles2}")

resps1 = sum(d['resps'] for d in data1)
resps2 = sum(d['resps'] for d in data2)
print(f"Total Responsibilities:  {resps1} -> {resps2} (+{resps2 - resps1})")

outc1 = sum(d['outcomes'] for d in data1)
outc2 = sum(d['outcomes'] for d in data2)
print(f"Total Outcomes:          {outc1} -> {outc2} (+{outc2 - outc1})")

facts1 = sum(d['facts'] for d in data1)
facts2 = sum(d['facts'] for d in data2)
print(f"Total Career Facts:      {facts1} -> {facts2} (+{facts2 - facts1})")

needs_rev1 = sum(1 for d in data1 if d['status'] == 'NEEDS_REVIEW')
needs_rev2 = sum(1 for d in data2 if d['status'] == 'NEEDS_REVIEW')
print(f"Resumes with NEEDS_REVIEW (Failing): {needs_rev1} (v1) -> {needs_rev2} (v2)")

partial1 = sum(1 for d in data1 if d['status'] == 'PARTIAL')
partial2 = sum(1 for d in data2 if d['status'] == 'PARTIAL')
print(f"Resumes with PARTIAL Grade:          {partial1} (v1) -> {partial2} (v2)")

good1 = sum(1 for d in data1 if d['status'] == 'GOOD')
good2 = sum(1 for d in data2 if d['status'] == 'GOOD')
print(f"Resumes with GOOD Grade:             {good1} (v1) -> {good2} (v2)")
