import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

json_path = r'd:\resume\Resume-Builder\scripts\batch_parsed_results.json'
with open(json_path, 'r', encoding='utf-8') as f:
    data = json.load(f)

wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

# Colors & Styles
NAVY_HEADER = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
REGULAR_FONT = Font(name='Calibri', size=10)
BOLD_FONT = Font(name='Calibri', size=10, bold=True)
TITLE_FONT = Font(name='Calibri', size=14, bold=True, color='1E293B')
SUBTITLE_FONT = Font(name='Calibri', size=10, italic=True, color='64748B')

THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

FILL_EXCELLENT = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid') # Green
FONT_EXCELLENT = Font(name='Calibri', size=10, bold=True, color='166534')

FILL_GOOD = PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid') # Blue
FONT_GOOD = Font(name='Calibri', size=10, bold=True, color='075985')

FILL_PARTIAL = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid') # Amber
FONT_PARTIAL = Font(name='Calibri', size=10, bold=True, color='92400E')

FILL_REVIEW = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid') # Rose
FONT_REVIEW = Font(name='Calibri', size=10, bold=True, color='991B1B')

def style_table(ws, headers, rows, title=''):
    curr_row = 1
    if title:
        ws.cell(row=curr_row, column=1, value=title).font = TITLE_FONT
        curr_row += 1
        ws.cell(row=curr_row, column=1, value=f'UpRole Enterprise Resume Extraction Audit | Generated: 2026-10-23 | Total Records: {len(rows)}').font = SUBTITLE_FONT
        curr_row += 2

    # Write headers
    header_row = curr_row
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[header_row].height = 28

    # Write rows
    for r_idx, row_data in enumerate(rows, start=header_row + 1):
        ws.row_dimensions[r_idx].height = 20
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            
            # Alignments
            if isinstance(val, (int, float)) or (isinstance(val, str) and val.endswith('%')):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

            # Badge styling for status
            if val in ['EXCELLENT', 'GOOD', 'PARTIAL', 'NEEDS_REVIEW']:
                if val == 'EXCELLENT':
                    cell.fill = FILL_EXCELLENT
                    cell.font = FONT_EXCELLENT
                elif val == 'GOOD':
                    cell.fill = FILL_GOOD
                    cell.font = FONT_GOOD
                elif val == 'PARTIAL':
                    cell.fill = FILL_PARTIAL
                    cell.font = FONT_PARTIAL
                else:
                    cell.fill = FILL_REVIEW
                    cell.font = FONT_REVIEW
                cell.alignment = Alignment(horizontal='center', vertical='center')

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < header_row:
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

# ── Sheet 1: Executive Summary ───────────────────────────────────────────
ws1 = wb.create_sheet(title='Executive_Summary')
headers1 = [
    'Resume ID', 'File Name', 'Format', 'File Size (KB)', 'Candidate Name',
    'Detected Headline', 'Stated Exp', 'Roles Found', 'Responsibilities',
    'Outcomes', 'Metrics', 'Skills Count', 'Education', 'Certs', 'Facts in Graph',
    'Accuracy Score', 'Quality Grade'
]
rows1 = []
for r in data['summaryRows']:
    rows1.append([
        r['resume_id'], r['file_name'], r['file_type'], float(r['file_size_kb']),
        r['candidate_name'], r['headline'], r['stated_experience'],
        r['roles_count'], r['responsibilities_count'], r['outcomes_count'],
        r['metrics_count'], r['skills_count'], r['education_count'],
        r['certifications_count'], r['career_facts_count'],
        str(r['accuracy_percentage']) + '%', r['quality_status']
    ])
style_table(ws1, headers1, rows1, 'UpRole 21-Resume Batch Extraction & Accuracy Comparison')

# ── Sheet 2: Identity Comparison ─────────────────────────────────────────
ws2 = wb.create_sheet(title='Identity_Comparison')
headers2 = [
    'Resume ID', 'Candidate Name', 'Professional Headline', 'Email Address',
    'Phone Number', 'Location', 'LinkedIn / URLs', 'Stated Experience', 'Professional Summary Snippet'
]
rows2 = []
for r in data['identityRows']:
    rows2.append([
        r['resume_id'], r['candidate_name'], r['headline'], r['email'],
        r['phone'], r['location'], r['linkedin'], r['stated_experience'],
        (r['summary'] or '')[:200]
    ])
style_table(ws2, headers2, rows2, 'Candidate Identity & Contact Extraction Breakdown')

# ── Sheet 3: Employment Roles ────────────────────────────────────────────
ws3 = wb.create_sheet(title='Employment_Roles')
headers3 = [
    'Resume ID', 'Candidate Name', 'Role #', 'Company / Organization', 'Job Title',
    'Start Date', 'End Date', 'Is Current Role', 'Work Location', 'Team Size',
    'Cross-Functional Teams', 'Responsibilities', 'Outcomes', 'Metrics'
]
rows3 = []
for r in data['rolesRows']:
    rows3.append([
        r['resume_id'], r['candidate_name'], r['role_number'], r['company'],
        r['job_title'], r['start_date'], r['end_date'],
        'Yes' if r['is_current_role'] else 'No', r['work_location'],
        r['team_size'], r['cross_functional_teams'],
        r['responsibilities_count'], r['outcomes_count'], r['metrics_count']
    ])
style_table(ws3, headers3, rows3, 'Extracted Employment Roles & Career Engagements')

# ── Sheet 4: Responsibilities vs Outcomes ────────────────────────────────
ws4 = wb.create_sheet(title='Responsibilities_vs_Outcomes')
headers4 = [
    'Resume ID', 'Candidate Name', 'Company', 'Job Title', 'Statement Type', 'Verbatim Bullet / Statement Text'
]
rows4 = []
for r in data['bulletRows'][:500]: # Cap first 500 bullets for performance
    rows4.append([
        r['resume_id'], r['candidate_name'], r['company'], r['job_title'],
        r['type'], r['bullet_text']
    ])
style_table(ws4, headers4, rows4, 'Operational Responsibilities vs Measurable Outcomes Classification (QA010)')

# ── Sheet 5: Metrics Breakdown ───────────────────────────────────────────
ws5 = wb.create_sheet(title='Metrics_Breakdown')
headers5 = [
    'Resume ID', 'Candidate Name', 'Company / Context', 'Metric Name', 'Extracted Value',
    'Unit', 'Baseline', 'Result', 'Original Verbatim Text'
]
rows5 = []
for r in data['metricsRows']:
    rows5.append([
        r['resume_id'], r['candidate_name'], r['company'], r['metric_name'],
        r['value'], r['unit'], r['baseline'], r['result'], r['original_text']
    ])
style_table(ws5, headers5, rows5, 'Granular Metrics Extraction (Scale, Percent, Efficiency - QA008 & QA009)')

# ── Sheet 6: Skills Inventory ────────────────────────────────────────────
ws6 = wb.create_sheet(title='Skills_Inventory')
headers6 = [
    'Resume ID', 'Candidate Name', 'Skill Name', 'Category', 'Proficiency Status'
]
rows6 = []
for r in data['skillsRows'][:600]:
    rows6.append([
        r['resume_id'], r['candidate_name'], r['skill_name'], r['category'], r['proficiency']
    ])
style_table(ws6, headers6, rows6, 'Skills & Tools Inventory (QA011 Rule: Proficiency Null Unless Stated)')

# ── Sheet 7: Education & Certifications ──────────────────────────────────
ws7 = wb.create_sheet(title='Education_&_Certifications')
headers7 = [
    'Resume ID', 'Candidate Name', 'Degree / Credential Name', 'Institution / Issuing Body', 'Completion Year'
]
rows7 = []
for r in data['educationRows']:
    rows7.append([
        r['resume_id'], r['candidate_name'], r['degree'], r['institution'], r['completion_year']
    ])
style_table(ws7, headers7, rows7, 'Qualifications, Degrees, and Professional Certifications')

# ── Sheet 8: QA & Audit Gaps ─────────────────────────────────────────────
ws8 = wb.create_sheet(title='QA_&_Audit_Gaps')
headers8 = [
    'Resume ID', 'Candidate Name', 'Accuracy Score', 'Extraction Grade',
    'Key Extraction Strengths', 'Identified Gaps & Layout Warnings'
]
rows8 = []
for r in data['auditRows']:
    rows8.append([
        r['resume_id'], r['candidate_name'], r['accuracy_score'], r['status'],
        r['strengths'], r['warnings']
    ])
style_table(ws8, headers8, rows8, 'Quality Assurance, Information Gaps & Parsing Audit')

paths = [
    r'd:\resume\UpRole_21_Resume_Fine_Tuned_Report_v3.xlsx',
    r'd:\resume\UpRole_21_Resume_Fine_Tuned_Report_v3_alt.xlsx',
    r'd:\resume\UpRole_21_Resume_Batch_Extraction_Report_v3.xlsx',
    r'd:\resume\UpRole_21_Resume_Batch_Extraction_Report_v2.xlsx'
]

for p in paths:
    try:
        wb.save(p)
        print(f'Excel Report successfully saved at: {p}')
        break
    except PermissionError:
        continue
